from flask import Blueprint, render_template, request, flash, url_for, jsonify, send_file
from flask_login import login_required, current_user
from .models import User, Schools, Inventar
from . import connectDB
from fpdf import FPDF, XPos, YPos
from docx import Document
import os, csv, openpyxl


create = Blueprint('create', __name__)

@create.route('/createFile', methods=['POST', 'GET'])
@login_required
def createFile():
    if request.method == 'POST':
        data = request.get_json()
        
        user_ids = data.get('userIds', [])
        invIds = data.get('invIds', [])
        page_type = data.get('pageType')

        try:
            with connectDB() as session:
                school_entry = session.query(Schools).filter_by(id=current_user.school_id).first()
                school_name = school_entry.short

                path = os.path.join('website', 'static', 'directories', school_name)

                try:
                    os.makedirs(path, exist_ok=True)
                except Exception as e:
                    print(f'Error: {e}')
        except Exception as e:
            print(f'Error: {e}')

        if user_ids:
            objects = getDataUser(user_ids)
            title = 'User List'
        elif invIds:
            objects = getDataInv(invIds)
            title = 'Inventar List'


        try:
            if page_type in ['print', 'pdf']:
                file_path = os.path.join('static', 'directories', school_name, f'{title}.pdf')

                try:
                    if os.path.exists(f'website/{file_path}'):
                        os.remove(f'website/{file_path}')
                            
                    createPDF(title, objects, path)
                except Exception as e:
                    print(f'Error: {e}') 
                    
                return jsonify({'file': file_path}), 200
            elif page_type == 'csv':
                file_path = os.path.join('static', 'directories', school_name, f'{title}.csv')

                try:
                    if os.path.exists(f'website/{file_path}'):
                        os.remove(f'website/{file_path}')
                            
                    createCSV(title, objects, path)
                except Exception as e:
                    print(f'Error: {e}') 
                        
                return jsonify({'file': file_path}), 200
            elif page_type == 'xlsx':
                file_path = os.path.join('static', 'directories', school_name, f'{title}.xlsx')

                try:
                    createXLSX(title, objects, path)
                except Exception as e:
                    print(f'Error: {e}') 
                        
                return jsonify({'file': file_path}), 200
            else:
                file_path = os.path.join('static', 'directories', school_name, f'{title}.docx')

                try:
                    if os.path.exists(f'website/{file_path}'):
                        os.remove(f'website/{file_path}')
                            
                    createDOCX(title, objects, path)
                except Exception as e:
                    print(f'Error: {e}') 
                        
                return jsonify({'file': file_path}), 200
        except Exception as e:
            print(f'Error: {e}')
            return jsonify({'Error': 'Error'}), 500
        

def getDataInv(inv_ids):
    with connectDB() as session:
        inventar = []
        for inv in inv_ids:
            inventory = session.query(Inventar).filter(
                Inventar.school_id == current_user.school_id,
                Inventar.uniqueid == inv
            ).first()
            user = session.query(User).filter_by(id = inventory.user_id).first()
            inv_data = {
                'ID UNIC': inv,
                'DENUMIRE': inventory.name,
                'DETINATOR': f'{user.nume} {user.prenume}',
                'CANTITATE': inventory.cant,
                'PRET/BUC': f'{inventory.price} RON',
                'VALOARE TOTALA': f'{inventory.total_price} RON' 
            }
            inventar.append(inv_data)
    return(inventar)


def getDataUser(user_ids):
    with connectDB() as session:
        users = []
        for user_id in user_ids:
            user_entry = session.query(User).filter(
                User.school_id == current_user.school_id,
                User.id == user_id
            ).first()
            user_data = {
                'ID': user_entry.id,
                'NUME': user_entry.nume,
                'PRENUME': user_entry.prenume,
                'ROL': user_entry.rol.capitalize(),
                'DEPARTMENT': user_entry.calificare.capitalize() if user_entry.rol != 'elev' else str(user_entry.clasa).upper()
            }
            users.append(user_data)
    return(users)


def createPDF(title, users, path):
    class PDF(FPDF):
        def header(self):
            self.add_font('Inter', '',
                        r"G:\My Drive\Management\website\fonts\Inter-Regular.ttf")
            
            self.add_font('Inter', 'B',
                        r'G:\My Drive\Management\website\fonts\Inter-Bold.ttf')
            
            self.set_font('Inter', 'B', 18)

            self.cell(0, 10, title, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align='C')
            self.ln()
        
        def footer(self):
            self.set_y(-15)

            self.set_font('Inter', '', 8)

            self.set_text_color(169, 169, 169)

            self.cell(0, 10, f'Page {self.page_no()}/{{nb}}', align='C')
    
    pdf = PDF('P', 'mm', 'A4')

    pdf.alias_nb_pages()

    pdf.set_auto_page_break(auto=True, margin=15)

    pdf.add_page()

    pdf.set_draw_color(50)
    pdf.set_line_width(.5)
    pdf.set_text_color(35, 50, 85)

    col_widths = [35, 40, 40, 35, 40]

    user_data = users[0]  # Choose an appropriate index based on your data structure

    # Add headers
    for key, width in zip(user_data.keys(), col_widths):
        pdf.set_fill_color(248, 248, 248)
        pdf.set_text_color(35, 50, 85)
        pdf.set_font("Inter", 'B', size=12)
        pdf.cell(width, 10, key, border='B', align='C', fill=True)

    pdf.ln()

    # Add data rows
    for user_data in users:
        pdf.set_fill_color(255, 255, 255)
        pdf.set_text_color(35, 50, 85)
        pdf.set_font("Inter", '', size=10)
        
        # Assuming keys in user_data are consistent across all items in the list
        for key, width in zip(user_data.keys(), col_widths):
            pdf.cell(width, 10, str(user_data[key]), border=0, align='L', fill=True)

        pdf.ln()

    pdf.output(f"{path}\\{title}.pdf")


def createDOCX(title, data, path):
    document = Document()

    field_names = list(data[0].keys())

    table = document.add_table(rows=1, cols=len(field_names))
    header_row = table.rows[0].cells
    for i, field_names in enumerate(field_names):
        header_row[i].text = field_names
    
    for row_data in data:
        row_cells = table.add_row().cells
        for i, value in enumerate(row_data.values()):
            row_cells[i].text = str(value)
    file_name = f'{path}/{title}.docx'

    document.save(file_name)
    

def createCSV(title, data, path):
    file_name = f'{path}/{title}.csv'

    field_names = data[0].keys()

    with open(file_name, mode='w', newline='') as file:
        writer = csv.DictWriter(file, fieldnames=field_names)

        writer.writeheader()

        writer.writerows(data)


def createXLSX(title, data, path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active

    field_names = list(data[0].keys())
    sheet.append(field_names)

    for row_data in data:
        sheet.append(list(row_data.values()))
    
    file_name = f'{path}/{title}.xlsx'

    workbook.save(file_name)